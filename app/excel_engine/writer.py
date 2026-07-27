import openpyxl
from pathlib import Path
from app.core.config import settings 
from app.exceptions.base import ProcessingError
from app.logging.setup import logger

class ExcelWriter:
    def write_multiple(self, template_path: str, records: list[dict], output_path: str):
     try:   
        wb = openpyxl.load_workbook(template_path)
        ws = wb.active

        print(records)

        # Map header names to column indices
        col_map = {}
        for col in range(1, ws.max_column + 1):
            header = ws.cell(row=1, column=col).value
            if header:
                col_map[str(header).strip()] = col

        # Remove all data rows below header
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)

        # Write each record as a new row
        for i, record in enumerate(records):
            row_num = i + 2
            for field_name, value in record.items():
                for header, col_idx in col_map.items():
                    if field_name.strip().lower() == header.lower():
                        ws.cell(row=row_num, column=col_idx, value=value if value else "")
                        break

        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        wb.close()
     except Exception as e:
        logger.error(f"Excel writing failed: {e}")
        raise ProcessingError("Failed to generate Excel file.") from e
