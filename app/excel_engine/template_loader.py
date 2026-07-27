import csv
import openpyxl
from pathlib import Path

class TemplateLoader:
    def __init__(self, header_row: int = 1, dummy_row: int = 2):
        self.header_row = header_row
        self.dummy_row = dummy_row

    def load_schema(self, template_path: str) -> list[dict]:
        """
        Returns a list of field definitions:
        [{"field_name": "...", "dummy_value": "..."}, ...]
        Supports .xlsx and .csv files.
        """
        ext = Path(template_path).suffix.lower()
        if ext == '.csv':
            return self._load_csv(template_path)
        elif ext == '.xlsx':
            return self._load_xlsx(template_path)
        else:
            raise ValueError(f"Unsupported template format '{ext}'. Please use .xlsx or .csv.")

    def _load_csv(self, path: str) -> list[dict]:
        with open(path, newline='', encoding='utf-8') as f:
            reader = csv.reader(f)
            rows = list(reader)
        # Row numbers are 1‑based
        headers = rows[self.header_row - 1] if len(rows) >= self.header_row else []
        dummies = rows[self.dummy_row - 1] if len(rows) >= self.dummy_row else []
        schema = []
        for i, h in enumerate(headers):
            dummy = dummies[i] if i < len(dummies) else ""
            schema.append({"field_name": h.strip(), "dummy_value": dummy.strip()})
        return schema

    def _load_xlsx(self, path: str) -> list[dict]:
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        headers = []
        dummies = []
        for col in range(1, ws.max_column + 1):
            header = ws.cell(row=self.header_row, column=col).value
            dummy = ws.cell(row=self.dummy_row, column=col).value
            if header is not None:
                headers.append(str(header).strip())
                dummies.append(str(dummy).strip() if dummy else "")
        wb.close()
        schema = []
        for name, dummy in zip(headers, dummies):
            schema.append({"field_name": name, "dummy_value": dummy})
        return schema